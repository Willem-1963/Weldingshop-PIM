#!/usr/bin/env python3
"""Prepare and validate the Rhodius 2026 Excel price list in a PIM test database."""

from __future__ import annotations

import csv
import json
import shutil
import sqlite3
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
EXCEL = ROOT.parent / "gesprekken/Tijdelijke RHODIUS Prijslijst 2026 NL Frezen Weldingshop.xlsx"
CANTO = ROOT / "output/rhodius_all_albums/products.json"
LIVE_DB = ROOT / "data/database/suppliers/rhodius-abrasives-gmbh.sqlite"
OUT = ROOT / "output/rhodius_excel_2026_preparation"
TEST_DB = OUT / "rhodius_import_test.sqlite"


def text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def number(value):
    if value in (None, ""):
        return None
    if not isinstance(value, str):
        return float(value)
    raw = value.strip().lower().replace(" kg", "")
    if not any(character.isdigit() for character in raw):
        return None
    if "," in raw:
        raw = raw.replace(".", "").replace(",", ".")
    return float(raw)


def ve_number(value):
    raw = text(value).lower().replace(" ", "")
    if "x" in raw:
        factors = [float(part.replace(",", ".")) for part in raw.split("x")]
        result = 1.0
        for factor in factors:
            result *= factor
        return result
    return float(raw)


def sales_unit(price_unit):
    return {"€/stuk": "stuk", "€/pak": "pak", "€/blik": "blik"}.get(text(price_unit), "stuk")


def read_excel():
    sheet = load_workbook(EXCEL, read_only=True, data_only=True)["Prijslijst 2026"]
    headers = [cell.value for cell in next(sheet.iter_rows())]
    return [
        dict(zip(headers, (cell.value for cell in row)))
        for row in sheet.iter_rows(min_row=2)
        if row[6].value not in (None, "")
    ]


def normalize(row):
    article = text(row["Artikelnummer"])
    unit = sales_unit(row["Prijseenheid"])
    ve = ve_number(row["VE"])
    piece_kg = number(row["Gewicht in kg/stuk"])
    package_kg = number(text(row["Gewicht/VE"]).replace(" kg", "")) if row["Gewicht/VE"] else None
    sales_kg = (
        round(piece_kg * ve, 9)
        if piece_kg is not None and unit in {"pak", "blik"}
        else piece_kg
    )
    gtin_piece = text(row["GTIN-code"])
    gtin_package = text(row["GTIN/verpakking"])
    leading_ean = gtin_package if unit in {"pak", "blik"} and gtin_package else gtin_piece
    title = " ".join(part for part in ["Rhodius", text(row["Productnaam"]), text(row["Afmetingen"]), text(row["Korrel / Draad Ø"])] if part)
    raw = dict(row)
    raw.update({
        "packaging_quantity_normalized": ve,
        "sales_unit_normalized": unit,
        "gtin_piece": gtin_piece,
        "gtin_package": gtin_package,
        "ean_selection_rule": "package_gtin_for_pack_or_tin_else_piece_gtin",
    })
    return {
        "sku": article,
        "supplier_sku": article,
        "ean": leading_ean,
        "vendor": "Rhodius Abrasives GmbH",
        "brand": "Rhodius",
        "source_title": title,
        "source_description": text(row["Opmerkingen"]),
        "price": number(row["Bruto prijs 2026"]),
        "cost_price": number(row["Netto prijs"]),
        "purchase_unit": unit,
        "sales_unit": unit,
        "purchase_units_per_sales_unit": 1.0,
        "unit_calculation_mode": "multiply",
        "purchase_discount_percent": number(row["Korting"]) * 100,
        "kg_per_purchase_unit": sales_kg,
        "kg_per_sales_unit": sales_kg,
        "weight_grams": round(sales_kg * 1000, 6) if sales_kg is not None else None,
        "stock_quantity": 0,
        "available": 0,
        "product_type": text(row["Productsoort"]),
        "category": text(row["Prijscategorie"]),
        "category_full": text(row["Productsoort"]),
        "source_present": 1,
        "shopify_status": "draft",
        "inventory_policy": "continue",
        "raw_data_json": json.dumps(raw, ensure_ascii=False, default=str),
        "filter_values_json": json.dumps([
            value for value in [
                text(row["Kwaliteitsklasse"]), text(row["Afmetingen"]),
                text(row["Korrel / Draad Ø"]), f"VE {text(row['VE'])}",
            ] if value
        ], ensure_ascii=False),
        "execution": text(row["Toepassing"]),
        "package_quantity": ve,
        "package_weight_kg_source": package_kg,
        "gtin_piece": gtin_piece,
        "gtin_package": gtin_package,
    }


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    rows = read_excel()
    grouped = {}
    conflicts = []
    for row in rows:
        sku = text(row["Artikelnummer"])
        if sku in grouped:
            conflicts.append({"article_number": sku, "kept": grouped[sku], "conflicting": row})
            continue
        grouped[sku] = row
    normalized = [normalize(row) for row in grouped.values()]

    canto = json.loads(CANTO.read_text(encoding="utf-8"))
    canto_by_article = {text(row["article_number"]): row for row in canto if row.get("article_number")}
    canto_by_ean = {}
    for row in canto:
        for key in ("ean", "canto_ean"):
            if row.get(key):
                canto_by_ean[text(row[key])] = row

    matches = []
    images_by_sku = {}
    for product in normalized:
        match = canto_by_article.get(product["sku"])
        method = "article_number" if match else ""
        if not match:
            match = canto_by_ean.get(product["gtin_piece"]) or canto_by_ean.get(product["gtin_package"])
            method = "gtin" if match else "unmatched"
        image_files = match["image_files"].split(" | ") if match else []
        images_by_sku[product["sku"]] = image_files
        matches.append({
            "article_number": product["sku"], "ean": product["ean"], "method": method,
            "canto_article_number": match["article_number"] if match else "",
            "image_count": len(image_files), "image_files": " | ".join(image_files),
        })

    shutil.copy2(LIVE_DB, TEST_DB)
    db = sqlite3.connect(TEST_DB)
    db.execute("PRAGMA foreign_keys=ON")
    db.execute("DELETE FROM product_images")
    db.execute("DELETE FROM products")
    now = datetime.now(timezone.utc).isoformat()
    db_columns = {row[1] for row in db.execute("PRAGMA table_info(products)")}
    insert_fields = [key for key in normalized[0] if key in db_columns]
    insert_fields += ["first_seen_at", "last_seen_at", "updated_at"]
    sql = f"INSERT INTO products({','.join(insert_fields)}) VALUES ({','.join('?' for _ in insert_fields)})"
    for product in normalized:
        values = [product.get(key, now if key in {"first_seen_at", "last_seen_at", "updated_at"} else None) for key in insert_fields]
        db.execute(sql, values)
        for position, image_file in enumerate(images_by_sku[product["sku"]], start=1):
            image_path = str((ROOT / "output/rhodius_all_albums" / image_file).resolve())
            db.execute(
                "INSERT OR IGNORE INTO product_images(sku,image_url,position,alt_text) VALUES (?,?,?,?)",
                (product["sku"], image_path, position, product["source_title"]),
            )
    db.commit()

    validation = {
        "excel_rows": len(rows),
        "unique_articles": len(normalized),
        "duplicate_article_conflicts": len(conflicts),
        "test_db_products": db.execute("SELECT COUNT(*) FROM products").fetchone()[0],
        "test_db_images": db.execute("SELECT COUNT(*) FROM product_images").fetchone()[0],
        "products_with_images": db.execute("SELECT COUNT(DISTINCT sku) FROM product_images").fetchone()[0],
        "products_without_images": db.execute("SELECT COUNT(*) FROM products p WHERE NOT EXISTS (SELECT 1 FROM product_images i WHERE i.sku=p.sku)").fetchone()[0],
        "null_cost_prices": db.execute("SELECT COUNT(*) FROM products WHERE cost_price IS NULL").fetchone()[0],
        "null_packaging": sum(product["package_quantity"] is None for product in normalized),
        "unit_distribution": dict(Counter(product["sales_unit"] for product in normalized)),
        "match_distribution": dict(Counter(row["method"] for row in matches)),
    }
    db.close()

    (OUT / "normalized_products.json").write_text(json.dumps(normalized, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "image_matches.json").write_text(json.dumps(matches, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "conflicts.json").write_text(json.dumps(conflicts, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "validation.json").write_text(json.dumps(validation, ensure_ascii=False, indent=2), encoding="utf-8")
    with (OUT / "image_matches.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=matches[0].keys(), delimiter=";")
        writer.writeheader()
        writer.writerows(matches)
    print(json.dumps(validation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
